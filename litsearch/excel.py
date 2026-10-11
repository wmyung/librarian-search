"""Search-strategy workbook in the medical library's layout, plus screening sheets."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .classify import FLAGS
from .project import db_info
from .sources import ctgov

THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="DDEBF7")
SUB = PatternFill("solid", fgColor="F2F2F2")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
BOLD = Font(bold=True)


def _cell(ws, ref, value, bold=False, fill=None, align=WRAP, border=True):
    c = ws[ref]
    c.value = value
    if bold:
        c.font = BOLD
    if fill:
        c.fill = fill
    c.alignment = align
    if border:
        c.border = BOX
    return c


def _query_text(db: str, q) -> str:
    return ctgov.describe(q) if db.startswith("ctgov") else str(q)


def build_workbook(path: Path, strategy: dict, ctx: dict) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Search strategy"
    dbs = ctx["db_order"]
    names = [db_info(d)["name"] for d in dbs if d != "citation"]
    _cell(ws, "A1", "Databases: " + ", ".join(names), bold=True, border=False, align=Alignment(vertical="top"))
    _cell(ws, "F1", "Search date", bold=True, fill=SUB)
    dates = sorted({e.get("search_date") or ctx["search_date"] for e in strategy.get("databases", [])
                    if e["id"] in ctx["retrieved"]} | ({ctx["citation_date"]} if ctx.get("citation_date") else set()))
    _cell(ws, "G1", " ~ ".join([dates[0], dates[-1]]) if len(dates) > 1 else ctx["search_date"])

    for col, h in zip("ABCD", ["DB", "Search", "Query", "Results"]):
        _cell(ws, f"{col}3", h, bold=True, fill=HEAD, align=CENTER)

    # summary table (F3:...)
    _cell(ws, "F3", "DBs", bold=True, fill=HEAD, align=CENTER)
    _cell(ws, "F4", "Retrieved", bold=True, fill=SUB, align=CENTER)
    _cell(ws, "F5", "After dedup", bold=True, fill=SUB, align=CENTER)
    col = 7
    for d in dbs:
        L = get_column_letter(col)
        _cell(ws, f"{L}3", db_info(d)["name"] if d != "pubmed" else "MEDLINE (PubMed)", bold=True, fill=HEAD, align=CENTER)
        _cell(ws, f"{L}4", ctx["retrieved"].get(d, 0), align=CENTER)
        _cell(ws, f"{L}5", ctx["contributed"].get(d, 0), align=CENTER)
        ws.column_dimensions[L].width = 14
        col += 1
    L = get_column_letter(col)
    first, last = get_column_letter(7), get_column_letter(col - 1)
    _cell(ws, f"{L}3", "Total", bold=True, fill=HEAD, align=CENTER)
    _cell(ws, f"{L}4", f"=SUM({first}4:{last}4)", align=CENTER)
    _cell(ws, f"{L}5", f"=SUM({first}5:{last}5)", align=CENTER)
    ws.column_dimensions[L].width = 10
    _cell(ws, "F7", "* After dedup: records newly contributed by each database in priority order (" + " > ".join(db_info(d)["short"] for d in dbs) +
          ")", border=False, align=Alignment(wrap_text=False))

    row = 4
    for entry in strategy.get("databases", []):
        db = entry["id"]
        lines = entry.get("lines", [])
        if not lines or not any(l.get("q") for l in lines):
            continue
        fin = ctx["final_index"].get(db, len(lines) - 1)
        start = row
        for i, ln in enumerate(lines):
            counts = ctx["line_counts"].get(db, {})
            res = counts.get(i, ln.get("results"))
            _cell(ws, f"B{row}", ln.get("label") or (f"#{i + 1}" if db in ("embase", "cochrane", "wos", "scopus") else i + 1), align=CENTER)
            _cell(ws, f"C{row}", _query_text(db, ln.get("q", "")), bold=(i == fin))
            _cell(ws, f"D{row}", res if res is not None else "", bold=(i == fin), align=CENTER)
            row += 1
        ws.merge_cells(f"A{start}:A{row - 1}")
        info = db_info(db)
        _cell(ws, f"A{start}", f"{info['name']}\n({info['platform']})", bold=True, align=CENTER)
        for r in range(start, row):
            ws[f"A{r}"].border = BOX
    if ctx.get("citation_log"):
        start = row
        for i, s in enumerate(ctx["citation_log"], 1):
            _cell(ws, f"B{row}", i, align=CENTER)
            txt = f"Seed: {s.get('seed')} — {s.get('title', s.get('error', ''))[:120]}  (references {s.get('backward', 0)}, cited-by {s.get('forward', 0)})"
            _cell(ws, f"C{row}", txt)
            _cell(ws, f"D{row}", s.get("backward", 0) + s.get("forward", 0), align=CENTER)
            row += 1
        ws.merge_cells(f"A{start}:A{row - 1}")
        _cell(ws, f"A{start}", "Citation searching\n(OpenAlex)", bold=True, align=CENTER)

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 80
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 3
    ws.column_dimensions["F"].width = 12
    ws.freeze_panes = "A4"

    _records_sheet(wb, ctx)
    _overlap_sheet(wb, ctx)
    if ctx.get("validation"):
        _validation_sheet(wb, ctx)
    _dups_sheet(wb, ctx)
    _prisma_sheet(wb, strategy, ctx)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _first_author(r):
    if not r["authors"]:
        return ""
    a = r["authors"][0].split(",")[0]
    return a + (" et al." if len(r["authors"]) > 2 else (" & " + r["authors"][1].split(",")[0] if len(r["authors"]) == 2 else ""))


def _records_sheet(wb, ctx):
    ws = wb.create_sheet("Records")
    flag_cols = [(k, lbl) for k, lbl, _t, _r in FLAGS]
    heads = ["No.", "Authors", "Year", "Title", "Journal/Source", "Vol(Issue):Pages", "DOI", "PMID", "Found in", "Type"] + \
            [lbl for _k, lbl in flag_cols] + ["Known relevant", "Previously screened", "Abstract",
                                               "Screening (include/exclude)", "Exclusion reason", "Notes"]
    for i, h in enumerate(heads, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = BOLD, HEAD, CENTER, BOX
    for n, r in enumerate(ctx["unique"], 2):
        vi = r["volume"] + (f"({r['issue']})" if r["issue"] else "") + (f":{r['pages']}" if r["pages"] else "")
        fl = ctx["flags"][r["uid"]]
        vals = [r["uid"], _first_author(r), r["year"], r["title"], r["journal"], vi, r["doi"], r["pmid"],
                ", ".join(db_info(d)["short"] for d in r["found_in"]), r["type"]] + \
               ["●" if fl.get(k) else "" for k, _l in flag_cols] + \
               ["●" if r["uid"] in ctx.get("known_uids", set()) else "",
                "●" if r["uid"] in ctx.get("previous_uids", set()) else "", r["abstract"][:32000], "", "", ""]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=n, column=i, value=v)
            c.alignment = Alignment(vertical="top", wrap_text=(i in (4,)))
        if r["doi"]:
            ws.cell(row=n, column=7).hyperlink = f"https://doi.org/{r['doi']}"
        if r["pmid"]:
            ws.cell(row=n, column=8).hyperlink = f"https://pubmed.ncbi.nlm.nih.gov/{r['pmid']}/"
    widths = [6, 18, 6, 60, 28, 14, 24, 10, 18, 9] + [9] * len(flag_cols) + [9, 10, 40, 14, 18, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{len(ctx['unique']) + 1}"


def _dups_sheet(wb, ctx):
    ws = wb.create_sheet("Duplicates")
    heads = ["No. (Trash)", "DB", "Title", "Year", "DOI", "PMID", "Kept record No.", "Match basis"]
    for i, h in enumerate(heads, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = BOLD, HEAD, CENTER, BOX
    for n, d in enumerate(ctx["dups"], 2):
        for i, v in enumerate([d["uid"], db_info(d["db"])["short"], d["title"], d["year"], d["doi"], d["pmid"],
                               d["dup_of"], d.get("dup_reason", "")], 1):
            ws.cell(row=n, column=i, value=v)
    for i, w in enumerate([10, 12, 70, 6, 26, 10, 12, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{len(ctx['dups']) + 1}"


def _prisma_sheet(wb, strategy, ctx):
    ws = wb.create_sheet("PRISMA-S")
    rows = ctx["prisma_s"]
    for i, h in enumerate(["#", "PRISMA-S item", "This search"], 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = BOLD, HEAD, CENTER, BOX
    for n, (num, item, desc) in enumerate(rows, 2):
        for i, v in enumerate([num, item, desc], 1):
            c = ws.cell(row=n, column=i, value=v)
            c.alignment, c.border = WRAP, BOX
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 100


def _head(ws, row, values):
    for i, h in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = BOLD, HEAD, CENTER, BOX


def _overlap_sheet(wb, ctx):
    """Source contribution and overlap (adapted from CiteSource)."""
    ws = wb.create_sheet("Source overlap")
    ov, dbs = ctx["overlap"], ctx["db_order"]
    _head(ws, 1, ["Source", "Retrieved", "Unique records found", "Found only in this source", "Unique contribution (priority order)"])
    for n, d in enumerate(dbs, 2):
        for i, v in enumerate([db_info(d)["name"], ctx["retrieved"].get(d, 0), ov["in_source"][d],
                               ov["only_in_source"][d], ctx["contributed"].get(d, 0)], 1):
            ws.cell(row=n, column=i, value=v).border = BOX
    r0 = len(dbs) + 4
    ws.cell(row=r0 - 1, column=1, value="Pairwise overlap (unique records found in both sources)").font = BOLD
    _head(ws, r0, [""] + [db_info(d)["short"] for d in dbs])
    for n, a in enumerate(dbs, r0 + 1):
        ws.cell(row=n, column=1, value=db_info(a)["short"]).font = BOLD
        for i, b in enumerate(dbs, 2):
            c = ws.cell(row=n, column=i, value=ov["matrix"][a][b])
            c.border = BOX
            if a == b:
                c.fill = SUB
    ws.column_dimensions["A"].width = 32
    for i in range(2, len(dbs) + 3):
        ws.column_dimensions[get_column_letter(i)].width = 16


def _validation_sheet(wb, ctx):
    ws = wb.create_sheet("Validation")
    v = ctx["validation"]
    found = sum(1 for k in v if k["found"])
    ws.cell(row=1, column=1, value=f"Known relevant studies retrieved: {found} of {len(v)} "
                                   f"(recall {found / len(v):.0%})").font = BOLD
    _head(ws, 3, ["Known study", "Found", "Sources", "Record No.", "Title", "In PubMed", "PubMed lines not matched"])
    for n, k in enumerate(v, 4):
        diag = k.get("diagnosis") or {}
        vals = [k["id"], "yes" if k["found"] else "NO", ", ".join(db_info(d)["short"] for d in k["sources"]),
                k.get("uid"), k.get("title", ""), {True: "yes", False: "no"}.get(diag.get("in_pubmed"), ""),
                ", ".join(f"#{x}" for x in diag.get("failed_lines", []))]
        for i, val in enumerate(vals, 1):
            ws.cell(row=n, column=i, value=val).border = BOX
    for col, w in zip("ABCDEFG", [26, 8, 22, 10, 70, 10, 24]):
        ws.column_dimensions[col].width = w
