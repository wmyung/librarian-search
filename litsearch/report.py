"""Statistics, PRISMA-S items, methods paragraph, reply-letter draft, RIS export and md->docx."""
from __future__ import annotations

import re
from pathlib import Path

from .classify import FLAGS
from .project import CITATION_SOURCES, db_info


def fmt_date_en(d: str) -> str:
    import datetime as dt
    try:
        x = dt.date.fromisoformat(d)
        return x.strftime("%-d %B %Y")
    except ValueError:
        return d


def db_dates(strategy: dict, ctx: dict) -> dict:
    out = {}
    for e in strategy.get("databases", []):
        out[e["id"]] = e.get("search_date") or ctx["search_date"]
    return out


def citation_runs(strategy: dict, ctx: dict) -> list[dict]:
    """Citation searching per platform: OpenAlex (LS cite) and exports imported from Scopus / Web of Science."""
    runs = []
    dates = db_dates(strategy, ctx)
    seeds = [str(s) for s in (strategy.get("citation") or {}).get("seeds") or []]
    if ctx.get("citation_log"):
        runs.append({"platform": "OpenAlex", "seeds": len([s for s in ctx["citation_log"] if not s.get("error")]),
                     "seed_ids": [], "date": ctx.get("citation_date", "")})
    for d, plat in (("citation_scopus", "Scopus"), ("citation_wos", "Web of Science Core Collection")):
        if d in ctx["db_order"]:
            runs.append({"platform": plat, "seeds": len(seeds), "seed_ids": seeds, "date": dates.get(d, ctx["search_date"])})
    return runs


def _cit_sentence(r: dict) -> str:
    n = f"{r['seeds']} key article{'s' if r['seeds'] != 1 else ''}" if r["seeds"] else "key articles"
    ids = f" (seeds: {', '.join(r['seed_ids'])})" if r.get("seed_ids") else ""
    return f"backward (reference lists) and forward (citing articles) citation searching of {n}{ids} in {r['platform']} ({r['date']})"


def prisma_s(strategy: dict, ctx: dict) -> list[tuple]:
    dbs = [d for d in ctx["db_order"] if d not in ("ctgov", "ictrp") + CITATION_SOURCES]
    regs = [d for d in ctx["db_order"] if d in ("ctgov", "ictrp")]
    dates = db_dates(strategy, ctx)
    cit = citation_runs(strategy, ctx)
    order = " > ".join(db_info(d)["short"] for d in ctx["db_order"])
    return [
        (1, "Database name", "; ".join(f"{db_info(d)['name']} ({db_info(d)['platform']})" for d in dbs)),
        (2, "Multi-database searching", "Each database was searched separately (no simultaneous multi-database searching)."),
        (3, "Study registries", "; ".join(db_info(d)["platform"] for d in regs) or "Not searched"),
        (4, "Online resources and browsing", strategy.get("other_sources", "Not applicable")),
        (5, "Citation searching", "; ".join(_cit_sentence(r)[0].upper() + _cit_sentence(r)[1:] for r in cit) if cit
         else "Not performed (recommended once key trials and systematic reviews are identified)"),
        (6, "Contacts", strategy.get("contacts", "Not applicable")),
        (7, "Other methods", strategy.get("other_methods") or (
            f"Validation set of {len(ctx['validation'])} known relevant studies used to check sensitivity "
            f"({sum(1 for k in ctx['validation'] if k['found'])} retrieved)" if ctx.get("validation") else "Not applicable")),
        (8, "Full search strategies", "All strategies are in the 'Search strategy' sheet of this workbook (suitable as a supplementary table)"),
        (9, "Limits and restrictions", strategy.get("limits_en") or strategy.get("limits", "")),
        (10, "Search filters", strategy.get("filters", "No study-design search filters were used")),
        (11, "Prior work", strategy.get("prior_work", "Not applicable")),
        (12, "Updates", strategy.get("updates") or (
            f"Update search: {len(ctx['previous_uids'])} records matched a previous search export and were marked as previously screened"
            if ctx.get("n_previous_records") else "Not applicable")),
        (13, "Dates of searches", "; ".join(f"{db_info(d)['short']}: {dates.get(d, ctx['search_date'])}" for d in ctx["db_order"] if d not in CITATION_SOURCES)),
        (14, "Peer review", strategy.get("peer_review", "No formal peer review of the search strategy (PRESS); librarian review recommended if required")),
        (15, "Total records", "; ".join(f"{db_info(d)['short']} {ctx['retrieved'].get(d, 0)}" for d in ctx["db_order"]) +
         f" (total {ctx['total_retrieved']})"),
        (16, "Deduplication", f"Automated deduplication on DOI or PMID match, or normalised title + publication year (+/-1) + first author; "
                              f"the record from the highest-priority source was kept ({order}). {ctx['n_dups']} duplicates removed, "
                              f"{ctx['n_unique']} unique records. Removed records remain in the EndNote library Trash for audit."),
    ]


def methods_en(strategy: dict, ctx: dict) -> str:
    dates = db_dates(strategy, ctx)
    dbs = [d for d in ctx["db_order"] if d not in ("ctgov", "ictrp", "europepmc") + CITATION_SOURCES]
    phr = {"pubmed": "MEDLINE (via PubMed)", "ovid_medline": "MEDLINE (via Ovid)", "embase": "Embase (via Embase.com)",
           "cochrane": "the Cochrane Library (Cochrane Database of Systematic Reviews and CENTRAL)",
           "wos": "Web of Science Core Collection", "scopus": "Scopus", "psycinfo": "APA PsycInfo", "cinahl": "CINAHL"}
    names = [phr.get(d, db_info(d)["name"]) for d in dbs]
    if "europepmc" in ctx["db_order"]:
        names.append("Europe PMC")
    lst = ", ".join(names[:-1]) + (", and " if len(names) > 2 else " and ") + names[-1] if len(names) > 1 else (names[0] if names else "")
    uniq_dates = sorted({dates.get(d, ctx["search_date"]) for d in ctx["db_order"] if d not in CITATION_SOURCES}) or [ctx["search_date"]]
    when = fmt_date_en(uniq_dates[0]) if len(uniq_dates) == 1 else "between " + fmt_date_en(uniq_dates[0]) + " and " + fmt_date_en(uniq_dates[-1])
    concepts = [c for c in strategy.get("concepts", []) if c.get("name")]
    cpt = "; ".join(f"{c['name']} ({', '.join(c.get('terms', [])[:12])})" for c in concepts)
    p = [f"A comprehensive literature search was conducted in {lst} from inception to {when}."] if lst else []
    if cpt:
        vocab = [v for d, v in (("pubmed", "MeSH"), ("ovid_medline", "MeSH"), ("embase", "Emtree"), ("cochrane", "MeSH"),
                                ("psycinfo", "APA Thesaurus"), ("cinahl", "CINAHL Headings")) if d in ctx["db_order"]]
        vocab = list(dict.fromkeys(vocab))
        vtxt = f"controlled vocabulary ({' and '.join(vocab)}) and " if vocab else ""
        p.append(f"Search strategies combined {vtxt}free-text terms for the following concepts: {cpt}.")
    p.append(strategy.get("limits_en") or strategy.get("limits") or "No restrictions on study design, language, or publication date were applied.")
    if any(d in ctx["db_order"] for d in ("ctgov", "ictrp")):
        regs = [("ClinicalTrials.gov" if d == "ctgov" else "the WHO ICTRP") for d in ctx["db_order"] if d in ("ctgov", "ictrp")]
        p.append(f"Ongoing and unpublished studies were sought in {' and '.join(regs)}.")
    for r in citation_runs(strategy, ctx):
        n = f"{r['seeds']} key article{'s' if r['seeds'] != 1 else ''}" if r["seeds"] else "key articles"
        ids = f" ({', '.join(r['seed_ids'])})" if r.get("seed_ids") else ""
        p.append(f"Backward and forward citation searching of {n}{ids} was performed in {r['platform']}"
                 f" on {fmt_date_en(r['date'] or ctx['search_date'])}.")
    if ctx.get("citation_log"):
        fa = (strategy.get("citation") or {}).get("filter_applied")
        if fa:
            p.append(f"Of the {fa['before']:,} citation records, {fa['after']:,} whose title, abstract, or keywords contained "
                     "terms for both key concepts were retained.")
    p.append(f"In total, {ctx['total_retrieved']:,} records were retrieved; after removal of {ctx['n_dups']:,} duplicates "
             f"(matching on DOI, PMID, and normalised title, publication year, and first author; EndNote 21), "
             f"{ctx['n_unique']:,} unique records remained for screening.")
    if ctx.get("validation"):
        v = ctx["validation"]
        f = sum(1 for k in v if k["found"])
        p.append(f"Search sensitivity was checked against a validation set of {len(v)} known relevant studies, "
                 f"of which {f} ({f / len(v):.0%}) were retrieved.")
    if ctx.get("n_previous_records"):
        p.append(f"Records already screened in the previous search ({ctx['n_previous_records']:,} records) were identified "
                 f"using the same deduplication rules; {ctx['n_unique'] - len(ctx['previous_uids']):,} records were new.")
    p.append("The full search strategies for each source are provided in Supplementary Table S1.")
    return " ".join(p)


def letter_draft(strategy: dict, ctx: dict) -> str:
    t = strategy.get("title")
    dates = db_dates(strategy, ctx)
    if ctx.get("citation_date"):
        dates["citation"] = ctx["citation_date"]
    multi = len({dates.get(d, ctx["search_date"]) for d in ctx["db_order"]}) > 1
    rows = "\n".join(f"| {db_info(d)['name']} | " + (f"{dates.get(d, ctx['search_date'])} | " if multi else "") +
                     f"{ctx['retrieved'].get(d, 0):,} | {ctx['contributed'].get(d, 0):,} |" for d in ctx["db_order"])
    head = ("| Source | Search date | Records retrieved | Unique contribution |\n|---|---|---:|---:|" if multi
            else "| Source | Records retrieved | Unique contribution |\n|---|---:|---:|")
    total = f"| **Total** | {'| ' if multi else ''}**{ctx['total_retrieved']:,}** | **{ctx['n_unique']:,}** |"
    when = "The results were as follows." if multi else f"The results as of {fmt_date_en(ctx['search_date'])} were as follows."
    flags = ", ".join(f"{lbl} {ctx['flag_counts'].get(k, 0)}" for k, lbl, _t, _r in FLAGS if ctx["flag_counts"].get(k))
    concepts = "\n".join(f"- **{c['name']}**: {', '.join(c.get('terms', []))}" for c in strategy.get("concepts", []) if c.get("name"))
    cit = citation_runs(strategy, ctx)
    n_cit = sum(ctx["contributed"].get(d, 0) for d in CITATION_SOURCES)
    cit_txt = (f"Reference lists and citing articles of key articles were checked ({'; '.join(r['platform'] for r in cit)}), "
               f"adding {n_cit} records not found by the database searches."
               if cit else "Once the key trials or systematic reviews are identified, we recommend checking their reference lists "
                           "and citing articles (citation searching).")
    sources = ", ".join(db_info(d)["name"] for d in ctx["db_order"] if d not in CITATION_SOURCES)
    prof_txt = ""
    if ctx.get("profile") == "systematic":
        extra = ctx.get("nonstandard_used") or []
        prof_txt = ("\nThis search used the systematic-review profile: bibliographic databases, trial registries and "
                    "citation indexes that can be reported under PRISMA-S" +
                    (f". By exception, {'; '.join(extra)} was also used; report it as a supplementary method." if extra
                     else " were searched (no open-index citation tools or AI search services)."))
        if ctx.get("systematic_gaps"):
            prof_txt += " Not yet covered: " + "; ".join(ctx["systematic_gaps"]) + "."
        prof_txt += "\n"
    val_txt = prev_txt = ""
    if ctx.get("validation"):
        v = ctx["validation"]
        f = sum(1 for k in v if k["found"])
        missed = [k["id"] for k in v if not k["found"]]
        val_txt = (f"\nTo check sensitivity, we tested the search against {len(v)} known relevant studies: {f} were retrieved"
                   + (f"; not retrieved: {', '.join(missed)} (see the Validation sheet for the reason)." if missed else ".") + "\n")
    if ctx.get("n_previous_records"):
        prev_txt = (f"\nCompared with your previous search, {len(ctx['previous_uids']):,} records were already screened and "
                    f"{ctx['n_unique'] - len(ctx['previous_uids']):,} are new (EndNote group \"New since previous search\").\n")
    return f"""Dear Researcher,

Thank you for your request regarding "{t}". We have completed the literature search and summarise it below.

## 1. Search strategy

{{{{2-3 sentences on how the scope was set in light of the review objective (e.g., no restriction by study design or outcome)}}}}

We searched {sources}, combining each database's controlled vocabulary with free-text terms for every concept.
{prof_txt}
{concepts}

## 2. Results and screening

{when}

{head}
{rows}
{total}

After deduplication, {ctx['n_unique']:,} unique records remain. To support screening, records are flagged by title, abstract and keywords ({flags}). These flags are screening aids, not search limits; every record is retained.

{{{{2-3 sentences of screening advice suited to the review type}}}}
{val_txt}{prev_txt}
## 3. Supplementary searching

{cit_txt}

## 4. Attachments

- Search strategy workbook (.xlsx): full line-by-line strategies with hit counts, record list for screening, deduplication log, PRISMA-S
- EndNote library (.zip): one group per source, smart groups (random*, systematic review, placebo, meta-analysis, etc.), duplicates kept in the Trash
- RIS file for other tools (Zotero, Covidence, Rayyan)
- Draft methods text for the manuscript: methods_PRISMA-S.md
- PRISMA 2020 flow diagram with the identification stage completed (PRISMA_2020_flow.svg)
- ASReview import file for AI-assisted screening (*_asreview.csv)

We hope this is helpful for your research.

Kind regards,
"""


# ---------------------------------------------------------------- RIS export
RIS_TY = {"JOUR": "JOUR", "BOOK": "BOOK", "CHAP": "CHAP", "CONF": "CONF", "THES": "THES", "RPRT": "RPRT",
          "TRIAL": "DBASE", "PREPRINT": "JOUR", "GEN": "GEN"}


def write_ris(path: Path, recs: list[dict]) -> None:
    out = []
    for r in recs:
        L = [f"TY  - {RIS_TY.get(r['type'], 'GEN')}", f"ID  - {r['uid']}"]
        L += [f"AU  - {a}" for a in r["authors"]]
        for tag, key in (("TI", "title"), ("T2", "journal"), ("PY", "year"), ("VL", "volume"), ("IS", "issue"),
                         ("AB", "abstract"), ("DO", "doi"), ("SN", "issn"), ("LA", "language"), ("AD", "address")):
            if r.get(key):
                L.append(f"{tag}  - {r[key]}")
        if r["pages"]:
            sp, _, ep = r["pages"].partition("-")
            L.append(f"SP  - {sp}")
            if ep:
                L.append(f"EP  - {ep}")
        L += [f"KW  - {k}" for k in dict.fromkeys(r["keywords"])]
        L += [f"UR  - {u}" for u in dict.fromkeys(r["urls"])]
        if r["pmid"]:
            L.append(f"AN  - {r['pmid']}")
        elif r["accession"]:
            L.append(f"AN  - {r['accession']}")
        L.append(f"DB  - {', '.join(db_info(d)['short'] for d in r.get('found_in', [r['db']]))}")
        L.append("ER  - ")
        out.append("\n".join(L))
    path.write_text("\n\n".join(out) + "\n", encoding="utf-8")


# ---------------------------------------------------------------- markdown -> docx
def md_to_docx(md: str, path: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)

    def add_runs(par, text):
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
            if part:
                run = par.add_run(part)
                run.bold = i % 2 == 1

    lines = md.split("\n")
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    tbl.append(cells)
                i += 1
            t = doc.add_table(rows=len(tbl), cols=max(len(r) for r in tbl))
            t.style = "Table Grid"
            for r, row in enumerate(tbl):
                for c, val in enumerate(row):
                    add_runs(t.cell(r, c).paragraphs[0], val)
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", ln)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
        elif re.match(r"^\s*[-*]\s+", ln):
            add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", ln))
        elif re.match(r"^\s*\d+\.\s+", ln):
            add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\.\s+", "", ln))
        elif ln.strip():
            add_runs(doc.add_paragraph(), ln)
        i += 1
    doc.save(path)
